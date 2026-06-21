import fitz
from docx import Document
from pptx import Presentation

def extract_pdf(path):
    records = []
    try:
        doc = fitz.open(path)
        for page_number, page in enumerate(doc,start=1):
            text = page.get_text()
            records.append({
                "page": page_number,
                "text": text
            })
        doc.close()
        combined = "".join(r["text"] for r in records).strip()
        if combined:
            return ("parsed", records)
        else:
            return ("image_only", records)
    except:
        return ("failed", records)


def extract_docx(path):
    records = []
    try:
        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs]
        full_text = "\n".join(paragraphs)
        return ("parsed", [{"page": 1, "text": full_text}])
    except:
        return ("failed", [])

def extract_pptx(path):
    records = []
    try:
        doc = Presentation(path)
        for page_number, slide in enumerate(doc.slides,start=1):
            slide_text = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    slide_text.append(shape.text)
            records.append({
                "page": page_number,
                "text": "\n".join(slide_text)
            })
        combined = "".join(r["text"] for r in records).strip()
        if combined:
            return ("parsed", records)
        else:
            return ("image_only", records)
    except:
        return ("failed", records)


def extract_file(path):
    if path.suffix.lower() == ".pdf":return extract_pdf(path)
    elif path.suffix.lower() == ".docx":return extract_docx(path)
    elif path.suffix.lower() == ".pptx":return extract_pptx(path)
    else: return ("skipped",[])