import fitz
from pathlib import Path
from docx import Document
import pandas as pd
from pptx import Presentation

def extract_pdf(path):
    records = []
    try:
        doc = fitz.open(path)
        for page_number, page in enumerate(doc,start=1):
            text = page.get_text()
            records.append({
                "page": page_number,
                "text": text
            })
        doc.close()
        combined = "".join(r["text"] for r in records).strip()
        if combined:
            return ("parsed", records)
        else:
            return ("image_only", records)
    except:
        return ("failed", records)


def extract_docx(path):
    records = []
    try:
        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs]
        full_text = "\n".join(paragraphs)
        return ("parsed", [{"page": 1, "text": full_text}])
    except:
        return ("failed", [])

def extract_pptx(path):
    records = []
    try:
        doc = Presentation(path)
        for page_number, slide in enumerate(doc.slides,start=1):
            slide_text = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    slide_text.append(shape.text)
            records.append({
                "page": page_number,
                "text": "\n".join(slide_text)
            })
        combined = "".join(r["text"] for r in records).strip()
        if combined:
            return ("parsed", records)
        else:
            return ("image_only", records)
    except:
        return ("failed", records)


def extract_file(path):
    if path.suffix.lower() == ".pdf":return extract_pdf(path)
    elif path.suffix.lower() == ".docx":return extract_docx(path)
    elif path.suffix.lower() == ".pptx":return extract_pptx(path)
    else: return ("skipped",[])

def build_index(corpus_root, inventory_csv, output_path):
    root = Path(corpus_root)
    df = pd.read_csv(inventory_csv)
    in_scope = df[df["scope"].isin(["deck","memo","summary"])]
    all_pages = []
    for idx, row in in_scope.iterrows():
        full_path = root / row["relative_path"]
        status, records = extract_file(full_path)
        for rec in records:
            all_pages.append({
                "document": row["filename"],
                "company": row["company"],
                "page": rec["page"],
                "text": rec["text"],
            })
        df.at[idx, "status"] = status
    df.to_csv(inventory_csv, index=False)
    pd.DataFrame(all_pages).to_json(output_path, orient="records", lines=True)


if __name__ == "__main__":
    build_index(
        corpus_root="data/raw/Space",
        inventory_csv="data/processed/inventory.csv",
        output_path="data/processed/extracted.jsonl",
    )
